"""Deterministic helpers for the unattended scheduled reports (CLAUDE.md rule 37).

The scheduled skills (/daily-watch, /weekly-scan, /rating-research) are run by
an LLM session with nobody at the keyboard. Everything that must not be left
to judgment lives here instead:

  preflight      prove the run can work BEFORE it starts — SEC and Yahoo are
                 reachable, the checkout is a clean, current `main`, GitHub is
                 logged in, no earlier scheduled run is still open — and take
                 a lock. Any failure exits 1 with a plain reason. A run that
                 cannot reach its data sources stops; it never degrades to a
                 partial report (the 2026-06..09 cloud scans did, for months).
  scan-filings   list recent 8-K / 6-K filings from SEC EDGAR for the watch
                 scope or the whole watchlist. Fetch failures are counted and
                 a high failure rate exits 1 — "no filings" and "could not
                 look" must never read the same.
  hold-tickets   park any trade ticket this run generated (rename to
                 held-ticket-*) so the executor's ticket-*.json glob skips it.
                 Scheduled runs re-score; they never release a trade.
  branch         create this run's branch (auto/<job>-<date>) from origin/main.
  log            append this run's one-line outcome to the committed run log
                 (tracking/scheduled-run-log.md). Every run logs, quiet or
                 not, so each one ends in a merged record and a missed run is
                 a missing date.
  land           stage ONLY research/score paths, refuse if anything else (or
                 anything under tracking/live/) is staged, commit, push the
                 run's own branch, open the PR.
  merge          wait for CI on that PR, squash-merge it, sync local main.
  release        drop the lock.

branch/land/merge exist so the unattended session needs no raw git or gh
write permission: the branch name, the staged paths and the PR merged are all
fixed here, not chosen by the model.

Names no order tool and sends nothing to the broker (rule 29). Imports
requests / yfinance / openpyxl lazily (CI minimal-env rule).
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import subprocess
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
LOCK_PATH = ROOT / ".cache" / "scheduled-run.lock"
LIVE_DIR = ROOT / "tracking" / "live"
WORKBOOK_LOCK = ROOT / "00-master" / "~$ai_supply_chain_scoring.xlsx"
LOCK_MAX_AGE_HOURS = 4          # a crashed run must not block the next one forever
RUN_BRANCH_PREFIX = "auto/"     # every scheduled run's branch: auto/<job>-<date>
MAX_FETCH_FAILURE_RATE = 0.20   # above this the scan is treated as blocked
FILING_FORMS = {"8-K", "8-K/A", "6-K"}
EARNINGS_ITEM = "2.02"
LAND_PATHS = ["per-stock", "00-master", "tracking"]   # the only stageable roots
PRIVATE_PREFIX = "tracking/live/"
RUN_LOG = ROOT / "tracking" / "scheduled-run-log.md"
RUN_LOG_HEADER = (
    "# Scheduled run log\n\n"
    "One line per scheduled run (rule 37), written by "
    "`scripts/scheduled_run.py log`. Append-only.\n\n")
CI_WAIT_SECONDS = 20 * 60


# ── pure decisions (unit-tested) ────────────────────────────────────────────

def lock_is_live(lock: dict | None, now: dt.datetime,
                 max_age_hours: float = LOCK_MAX_AGE_HOURS) -> bool:
    """A lock blocks a new run only while it is younger than max_age_hours."""
    if not lock or "started" not in lock:
        return False
    try:
        started = dt.datetime.fromisoformat(lock["started"])
    except ValueError:
        return False
    return (now - started) < dt.timedelta(hours=max_age_hours)


# Status files the launchd executor's morning reconciliation rewrites in the
# main checkout every trading day (already sanitized, rule 29). They are not
# someone's unfinished work: the run carries them onto its branch and `land`
# commits them. Blocking on them refused every afternoon run (2026-10-13).
MACHINE_WRITTEN = frozenset({"tracking/live-status.json",
                             "tracking/live-vs-model.json"})


def git_problems(branch: str, modified: list[str], ahead: int) -> list[str]:
    """Why this checkout cannot host an unattended run (empty = fine)."""
    out = []
    modified = [m for m in modified if m not in MACHINE_WRITTEN]
    if branch != "main":
        out.append(f"checkout is on '{branch}', not main — another session is "
                   f"mid-work here")
    if modified:
        shown = ", ".join(modified[:5]) + (" …" if len(modified) > 5 else "")
        out.append(f"{len(modified)} tracked file(s) have uncommitted changes "
                   f"({shown})")
    if ahead:
        out.append(f"local main is {ahead} commit(s) ahead of origin/main "
                   f"(unpushed work)")
    return out


def open_run_prs(prs: list[dict]) -> list[str]:
    """Open PRs left behind by an earlier scheduled run."""
    return [f"#{p['number']} ({p['headRefName']})" for p in prs
            if str(p.get("headRefName", "")).startswith(RUN_BRANCH_PREFIX)]


def filings_in_window(recent: dict, start: dt.date, end: dt.date,
                      forms: set[str] = FILING_FORMS) -> list[dict]:
    """Filter a submissions-JSON `filings.recent` block to forms in [start, end]."""
    out = []
    cols = [recent.get(k, []) for k in
            ("form", "accessionNumber", "filingDate", "primaryDocument", "items")]
    for form, acc, fdate, doc, items in zip(*cols):
        if form not in forms:
            continue
        try:
            d = dt.date.fromisoformat(fdate)
        except (TypeError, ValueError):
            continue
        if start <= d <= end:
            item_list = [i.strip() for i in str(items or "").split(",") if i.strip()]
            out.append({"form": form, "date": fdate, "accession": acc,
                        "doc": doc, "items": item_list,
                        "earnings": EARNINGS_ITEM in item_list})
    return out


def scan_blocked(n_attempted: int, n_failed: int,
                 max_rate: float = MAX_FETCH_FAILURE_RATE) -> bool:
    """True when too many lookups failed for the result to mean anything."""
    if n_attempted == 0:
        return True
    return (n_failed / n_attempted) > max_rate


def new_actionable_tickets(now_names: list[str], before_names: list[str]) -> list[str]:
    """Tickets the executor would pick up that did not exist at preflight."""
    before = set(before_names)
    return sorted(n for n in now_names
                  if n.startswith("ticket-") and n.endswith(".json")
                  and n not in before)


def staged_violations(paths: list[str]) -> list[str]:
    """Staged files a scheduled run must never commit (empty = fine)."""
    ok_roots = tuple(r + "/" for r in LAND_PATHS)
    return sorted(p for p in paths
                  if p.startswith(PRIVATE_PREFIX) or not p.startswith(ok_roots))


def run_branch(lock: dict | None) -> str | None:
    """The one branch this run may create, push and merge."""
    if not lock or not lock.get("job") or not lock.get("started"):
        return None
    job = "".join(c for c in str(lock["job"]) if c.isalnum() or c == "-")
    return f"{RUN_BRANCH_PREFIX}{job}-{str(lock['started'])[:10]}"


def run_log_line(lock: dict | None, text: str) -> str | None:
    """`- **<date>** (<job>) — <text>` for this run, or None without a lock.
    The text is flattened to one line so the file stays one line per run."""
    if not lock or not lock.get("job") or not lock.get("started"):
        return None
    flat = " ".join(str(text).split())
    if not flat:
        return None
    return f"- **{str(lock['started'])[:10]}** ({lock['job']}) — {flat}"


# ── side-effecting shell ────────────────────────────────────────────────────

def _run(cmd: list[str], timeout: int = 60) -> tuple[int, str]:
    try:
        p = subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True,
                           timeout=timeout)
        return p.returncode, (p.stdout + p.stderr).strip()
    except (OSError, subprocess.TimeoutExpired) as e:
        return 1, str(e)


def _ticket_names() -> list[str]:
    tdir = LIVE_DIR / "tickets"
    return sorted(p.name for p in tdir.glob("ticket-*.json")) if tdir.exists() else []


def _read_lock() -> dict | None:
    try:
        return json.loads(LOCK_PATH.read_text())
    except (OSError, ValueError):
        return None


def _check_sec() -> str | None:
    try:
        from common import sec_get
        r = sec_get("https://data.sec.gov/submissions/CIK0001045810.json",
                    host_data_sec=True)
        if "filings" not in r.json():
            return "SEC EDGAR answered but without filing data"
    except Exception as e:                       # noqa: BLE001 — report, don't trace
        return f"SEC EDGAR unreachable ({type(e).__name__}: {str(e)[:120]})"
    return None


def _check_yahoo() -> str | None:
    try:
        import yfinance as yf
        h = yf.Ticker("NVDA").history(period="5d")
        if not len(h):
            return "Yahoo Finance returned no price rows for NVDA"
    except Exception as e:                       # noqa: BLE001
        return f"Yahoo Finance unreachable ({type(e).__name__}: {str(e)[:120]})"
    return None


def preflight(job: str) -> int:
    now = dt.datetime.now()
    problems: list[str] = []

    lock = _read_lock()
    if lock_is_live(lock, now):
        problems.append(f"another scheduled run holds the lock "
                        f"({lock.get('job')} since {lock.get('started')})")

    if WORKBOOK_LOCK.exists():
        problems.append("the scoring workbook is open in Excel — close it")

    _, branch = _run(["git", "rev-parse", "--abbrev-ref", "HEAD"])
    _, status = _run(["git", "status", "--porcelain", "--untracked-files=no"])
    modified = [ln.split(maxsplit=1)[1] for ln in status.splitlines()
                if len(ln.split(maxsplit=1)) == 2]
    rc, out = _run(["git", "fetch", "origin", "main"], timeout=120)
    if rc:
        problems.append(f"git fetch failed ({out[:120]})")
    _, ahead = _run(["git", "rev-list", "--count", "origin/main..main"])
    problems += git_problems(branch, modified, int(ahead) if ahead.isdigit() else 0)

    rc, out = _run(["gh", "pr", "list", "--state", "open", "--json",
                    "number,headRefName"])
    if rc:
        problems.append(f"GitHub CLI not usable ({out[:120]}) — run `gh auth login`")
    else:
        stuck = open_run_prs(json.loads(out or "[]"))
        if stuck:
            problems.append("an earlier scheduled run's PR is still open: "
                            + ", ".join(stuck))

    for check in (_check_sec, _check_yahoo):
        msg = check()
        if msg:
            problems.append(msg)

    if not problems:                              # only now touch the checkout
        rc, out = _run(["git", "merge", "--ff-only", "origin/main"])
        if rc:
            problems.append(f"could not fast-forward main ({out[:120]})")

    if problems:
        print(json.dumps({"ok": False, "job": job, "problems": problems},
                         indent=2, ensure_ascii=False))
        return 1

    LOCK_PATH.parent.mkdir(parents=True, exist_ok=True)
    LOCK_PATH.write_text(json.dumps({
        "job": job, "started": now.isoformat(timespec="seconds"),
        "tickets_before": _ticket_names()}, indent=2))
    print(json.dumps({"ok": True, "job": job,
                      "branch": f"{RUN_BRANCH_PREFIX}{job}-{now:%Y-%m-%d}",
                      "date": f"{now:%Y-%m-%d}"}, indent=2))
    return 0


def _scope(kind: str) -> list[str]:
    if kind == "watch":
        import earnings_sentinel as es
        return es.build_scope(es._holdings(),
                              es.latest_ranked_rows(es.SCORE_HISTORY))
    from openpyxl import load_workbook
    ws = load_workbook(ROOT / "00-master" / "ai_supply_chain_scoring.xlsx",
                       read_only=True, data_only=True)["Watchlist"]
    out = []
    for (t,) in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        if not t:
            break
        out.append(str(t).strip().upper())
    return out


def scan_filings(scope: str, days: int, out_path: str | None) -> int:
    from common import cik_for, sec_get
    end = dt.date.today()
    start = end - dt.timedelta(days=days)
    tickers = _scope(scope)
    filings, no_cik, failed = [], [], []
    for t in tickers:
        if "." in t:                    # foreign local line: no EDGAR filings
            no_cik.append(t)
            continue
        cik = cik_for(t)
        if not cik:
            no_cik.append(t)
            continue
        try:
            data = sec_get(f"https://data.sec.gov/submissions/CIK{cik}.json",
                           host_data_sec=True).json()
        except Exception as e:                    # noqa: BLE001
            failed.append({"ticker": t, "error": f"{type(e).__name__}: {str(e)[:80]}"})
            continue
        for f in filings_in_window(data.get("filings", {}).get("recent", {}),
                                   start, end):
            acc = f["accession"].replace("-", "")
            f.update(ticker=t, url=(f"https://www.sec.gov/Archives/edgar/data/"
                                    f"{int(cik)}/{acc}/{f['doc']}"))
            filings.append(f)
        time.sleep(0.02)
    attempted = len(tickers) - len(no_cik)
    report = {"scope": scope, "from": start.isoformat(), "to": end.isoformat(),
              "tickers": len(tickers), "attempted": attempted,
              "filings": sorted(filings, key=lambda f: (f["date"], f["ticker"])),
              "earnings": sorted({f["ticker"] for f in filings if f["earnings"]}),
              "no_edgar": no_cik, "failed": failed,
              "blocked": scan_blocked(attempted, len(failed))}
    text = json.dumps(report, indent=2)
    if out_path:
        Path(out_path).write_text(text + "\n")
        print(f"{len(filings)} filings from {attempted} names "
              f"({len(failed)} lookups failed, {len(no_cik)} not on EDGAR) "
              f"→ {out_path}")
    else:
        print(text)
    if report["blocked"]:
        print("SCAN BLOCKED: too many SEC lookups failed — do not report this "
              "window as quiet", file=sys.stderr)
        return 1
    return 0


def hold_tickets() -> int:
    lock = _read_lock() or {}
    fresh = new_actionable_tickets(_ticket_names(), lock.get("tickets_before", []))
    tdir = LIVE_DIR / "tickets"
    held = []
    for name in fresh:
        target = tdir / f"held-{name}"
        n = 2
        while target.exists():                    # append-only, never overwrite
            target = tdir / f"held-{Path(name).stem}.{n}.json"
            n += 1
        (tdir / name).rename(target)
        held.append(target.name)
    print(json.dumps({"held": held}, indent=2))
    return 0


def _fail(msg: str) -> int:
    print(json.dumps({"ok": False, "problem": msg}, indent=2, ensure_ascii=False))
    return 1


def _current_branch() -> str:
    return _run(["git", "rev-parse", "--abbrev-ref", "HEAD"])[1]


def branch() -> int:
    name = run_branch(_read_lock())
    if not name:
        return _fail("no lock — run preflight first")
    rc, out = _run(["git", "checkout", "-b", name, "origin/main"])
    if rc:
        return _fail(f"could not create {name}: {out[:160]}")
    print(json.dumps({"ok": True, "branch": name}))
    return 0


def log(text: str) -> int:
    lock = _read_lock()
    name = run_branch(lock)
    if not name or _current_branch() != name:
        return _fail(f"not on this run's branch ({name}) — run branch first")
    line = run_log_line(lock, text)
    if not line:
        return _fail("nothing to log")
    body = RUN_LOG.read_text() if RUN_LOG.exists() else RUN_LOG_HEADER
    RUN_LOG.write_text(body.rstrip("\n") + "\n" + line + "\n")
    print(json.dumps({"ok": True, "logged": line}, ensure_ascii=False))
    return 0


def land(message: str, body_file: str | None) -> int:
    lock = _read_lock()
    name = run_branch(lock)
    if not name or _current_branch() != name:
        return _fail(f"not on this run's branch ({name})")
    _run(["git", "add", "--", *LAND_PATHS])
    _, staged = _run(["git", "diff", "--cached", "--name-only"])
    paths = [p for p in staged.splitlines() if p.strip()]
    bad = staged_violations(paths)
    if bad:
        _run(["git", "reset", "-q"])
        return _fail("refusing to commit: " + ", ".join(bad[:8]))
    if not paths:
        return _fail("nothing to commit")
    rc, out = _run(["git", "commit", "-q", "-m", message])
    if rc:
        return _fail(f"commit failed: {out[:160]}")
    rc, out = _run(["git", "push", "-u", "origin", f"HEAD:refs/heads/{name}"],
                   timeout=180)
    if rc:
        return _fail(f"push failed: {out[:160]}")
    cmd = ["gh", "pr", "create", "--head", name, "--base", "main",
           "--title", message.splitlines()[0][:120]]
    cmd += ["--body-file", body_file] if body_file else ["--body", message]
    rc, out = _run(cmd, timeout=120)
    if rc:
        return _fail(f"PR not created: {out[:160]}")
    rc, num = _run(["gh", "pr", "view", name, "--json", "number", "-q", ".number"])
    lock["pr"] = int(num) if num.isdigit() else None
    LOCK_PATH.write_text(json.dumps(lock, indent=2))
    print(json.dumps({"ok": True, "pr": lock["pr"], "files": len(paths)}))
    return 0


def merge() -> int:
    lock = _read_lock() or {}
    name, pr = run_branch(lock), lock.get("pr")
    if not name or not pr:
        return _fail("no PR recorded for this run — run land first")
    rc, head = _run(["gh", "pr", "view", str(pr), "--json", "headRefName",
                     "-q", ".headRefName"])
    if rc or head != name:
        return _fail(f"PR #{pr} is not this run's branch ({head!r} != {name!r})")
    deadline = time.time() + CI_WAIT_SECONDS
    while True:
        _, checks = _run(["gh", "pr", "checks", str(pr)])
        if "fail" in checks:
            return _fail(f"CI failed on PR #{pr} — left open")
        if "pass" in checks and "pending" not in checks:
            break
        if time.time() > deadline:
            return _fail(f"CI did not finish on PR #{pr} — left open")
        time.sleep(10)
    rc, out = _run(["gh", "pr", "merge", str(pr), "--squash", "--delete-branch"],
                   timeout=180)
    if rc:
        return _fail(f"merge failed: {out[:160]}")
    _run(["git", "checkout", "main"])
    _run(["git", "fetch", "origin", "main"], timeout=120)
    rc, out = _run(["git", "merge", "--ff-only", "origin/main"])
    if rc:
        return _fail(f"merged, but local main did not fast-forward: {out[:120]}")
    _run(["git", "branch", "-D", name])
    print(json.dumps({"ok": True, "merged": pr}))
    return 0


def release() -> int:
    LOCK_PATH.unlink(missing_ok=True)
    print("lock released")
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("preflight")
    p.add_argument("--job", required=True)
    s = sub.add_parser("scan-filings")
    s.add_argument("--scope", choices=["watch", "all"], default="watch")
    s.add_argument("--days", type=int, default=3)
    s.add_argument("--out")
    sub.add_parser("hold-tickets")
    sub.add_parser("branch")
    lg = sub.add_parser("log")
    lg.add_argument("--text", required=True)
    ln = sub.add_parser("land")
    ln.add_argument("--message", required=True)
    ln.add_argument("--body-file")
    sub.add_parser("merge")
    sub.add_parser("release")
    a = ap.parse_args(argv)
    if a.cmd == "preflight":
        return preflight(a.job)
    if a.cmd == "scan-filings":
        return scan_filings(a.scope, a.days, a.out)
    if a.cmd == "hold-tickets":
        return hold_tickets()
    if a.cmd == "branch":
        return branch()
    if a.cmd == "log":
        return log(a.text)
    if a.cmd == "land":
        return land(a.message, a.body_file)
    if a.cmd == "merge":
        return merge()
    return release()


if __name__ == "__main__":
    sys.exit(main())
