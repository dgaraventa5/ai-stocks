#!/usr/bin/env python3
"""Rule 39: every name in the model portfolio has a three-sentence thesis.

Lists Targets names (Include?=Y) whose thesis.md section 1 is missing, still
the template, or not exactly three sentences. Exit 1 if any.
"""
import re
import sys
from pathlib import Path

_REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_REPO_ROOT / 'scripts'))


def portfolio_tickers(root: Path = _REPO_ROOT) -> list[str]:
    from openpyxl import load_workbook   # lazy (minimal-CI convention)
    ws = load_workbook(root / '00-master' / 'portfolio.xlsx',
                       read_only=True)['Targets']
    rows = ws.iter_rows(min_row=2, values_only=True)
    col = list(next(rows)).index('Include?')
    return [r[0] for r in rows
            if r[0] and r[0] != 'TOTAL' and r[col] == 'Y']


def sentence_count(text: str) -> int:
    return len(re.split(r'(?<=[.!?])\s+', text.strip()))


def problems(root: Path = _REPO_ROOT) -> dict[str, str]:
    from export_site_data import export_theses
    tickers = portfolio_tickers(root)
    out = {}
    for t, text in export_theses(root, tickers).items():
        if text is None:
            out[t] = 'no thesis'
        elif sentence_count(text) != 3:
            out[t] = f'{sentence_count(text)} sentences, needs 3'
    return out


if __name__ == '__main__':
    found = problems()
    for t, why in sorted(found.items()):
        print(f'{t}: {why}')
    if not found:
        print('every portfolio name has a three-sentence thesis')
    sys.exit(1 if found else 0)
